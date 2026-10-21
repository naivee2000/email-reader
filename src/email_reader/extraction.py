from __future__ import annotations

import codecs
import csv
import io
import mimetypes
import re
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any


@dataclass(slots=True)
class ExtractionResult:
    text: str
    metadata: dict[str, Any] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


class AttachmentExtractor:
    # Bump when extraction behavior changes, so cached local results are rebuilt.
    CACHE_VERSION = 2
    TEXT_CACHE_VERSION = 3
    OFFICE_CACHE_VERSION = 3
    TEXT_SUFFIXES = {".txt", ".md", ".json", ".xml", ".csv", ".tsv"}
    TEXT_TYPES = {
        "text/plain",
        "text/csv",
        "text/tab-separated-values",
        "text/markdown",
        "application/json",
        "application/xml",
        "text/xml",
    }

    @classmethod
    def version_for(cls, filename: str, mime_type: str) -> int:
        suffix = Path(filename).suffix.lower()
        media_type = (mime_type or "").split(";", 1)[0].strip().lower()
        if suffix in {".docx", ".xlsx", ".xlsm", ".pptx"}:
            return cls.OFFICE_CACHE_VERSION
        return (cls.TEXT_CACHE_VERSION if suffix in cls.TEXT_SUFFIXES
                or media_type in cls.TEXT_TYPES else cls.CACHE_VERSION)

    @staticmethod
    def _decode_text(content: bytes) -> str:
        if content.startswith((codecs.BOM_UTF32_LE, codecs.BOM_UTF32_BE)):
            return content.decode("utf-32", errors="replace")
        if content.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
            return content.decode("utf-16", errors="replace")
        if content.startswith(codecs.BOM_UTF8):
            return content.decode("utf-8-sig", errors="replace")
        # Some mail exports omit the UTF-16 BOM. Detect only a strong NUL
        # pattern so ordinary UTF-8 text is never reinterpreted as UTF-16.
        sample = content[:4096]
        if len(sample) >= 8 and len(sample) % 2 == 0:
            even = sample[::2].count(0) / len(sample[::2])
            odd = sample[1::2].count(0) / len(sample[1::2])
            if odd > .35 and even < .05:
                return content.decode("utf-16-le", errors="replace")
            if even > .35 and odd < .05:
                return content.decode("utf-16-be", errors="replace")
        return content.decode("utf-8", errors="replace")

    def extract(self, filename: str, mime_type: str, content: bytes) -> ExtractionResult:
        suffix = Path(filename).suffix.lower()
        mime_type = (mime_type or mimetypes.guess_type(filename)[0]
                     or "application/octet-stream").split(";", 1)[0].strip().lower()
        try:
            if suffix in {".csv", ".tsv"} or mime_type in {
                "text/csv", "text/tab-separated-values"
            }:
                return self._csv(content, delimiter="\t" if suffix == ".tsv"
                                 or mime_type == "text/tab-separated-values" else ",")
            if mime_type in self.TEXT_TYPES or suffix in self.TEXT_SUFFIXES:
                return ExtractionResult(self._decode_text(content))
            if suffix == ".pdf" or mime_type == "application/pdf":
                return self._pdf(content)
            if suffix == ".docx":
                return self._docx(content)
            if suffix in {".xlsx", ".xlsm"}:
                return self._xlsx(content)
            if suffix == ".pptx":
                return self._pptx(content)
            if mime_type.startswith("image/") or suffix in {
                ".png",
                ".jpg",
                ".jpeg",
                ".tiff",
                ".heic",
            }:
                return self._image_ocr(content, suffix or ".png")
        except Exception as exc:
            return ExtractionResult("", warnings=[f"Extraction failed: {type(exc).__name__}"])
        return ExtractionResult(
            "", warnings=["Original preserved; full-text extraction is unsupported"]
        )

    def _csv(self, content: bytes, *, delimiter: str = ",") -> ExtractionResult:
        decoded = self._decode_text(content)
        rows = list(csv.reader(io.StringIO(decoded), delimiter=delimiter))
        text = "\n".join(" | ".join(cell for cell in row) for row in rows)
        return ExtractionResult(text, {"rows": len(rows)})

    def _pdf(self, content: bytes) -> ExtractionResult:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        pages: list[str] = []
        ocr_pages: list[int] = []
        warnings: list[str] = []
        for number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            needs_ocr = len(text) < 30
            if not needs_ocr:
                try:
                    needs_ocr = bool(page.images) and self._pdf_page_has_large_image(
                        content, number - 1
                    )
                except Exception:
                    # Native text remains usable if image inspection fails.
                    needs_ocr = False
            if needs_ocr:
                try:
                    ocr = self._ocr_pdf_page(content, number - 1)
                except Exception as exc:
                    ocr = ExtractionResult("", warnings=[f"OCR failed: {type(exc).__name__}"])
                warnings.extend(f"Page {number}: {warning}" for warning in ocr.warnings)
                if ocr.text:
                    if text:
                        native = re.sub(r"\s+", " ", text).casefold()
                        extra_lines = [line for line in ocr.text.splitlines()
                                       if re.sub(r"\s+", " ", line).strip().casefold()
                                       not in native]
                        if extra_lines:
                            text += "\n" + "\n".join(extra_lines)
                    else:
                        text = ocr.text
                    ocr_pages.append(number)
            pages.append(f"[Page {number}]\n{text}".strip())
        return ExtractionResult(
            "\n\n".join(pages),
            {"pages": len(reader.pages), "ocr_pages": ocr_pages},
            warnings,
        )

    @staticmethod
    def _pdf_page_has_large_image(content: bytes, page_index: int) -> bool:
        import fitz

        with fitz.open(stream=content, filetype="pdf") as document:
            page = document.load_page(page_index)
            page_area = page.rect.get_area()
            if page_area <= 0:
                return False
            image_area = sum(
                (rect & page.rect).get_area()
                for image in page.get_images(full=True)
                for rect in page.get_image_rects(image[0])
            )
            return image_area >= page_area * .25

    def _ocr_pdf_page(self, content: bytes, page_index: int) -> ExtractionResult:
        import fitz

        with fitz.open(stream=content, filetype="pdf") as document:
            page = document.load_page(page_index)
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            return self._image_ocr(pixmap.tobytes("png"), ".png")

    def _docx(self, content: bytes) -> ExtractionResult:
        from docx import Document

        document = Document(io.BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table_index, table in enumerate(document.tables, start=1):
            paragraphs.append(f"[Table {table_index}]")
            paragraphs.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        # Policy numbers, contract dates, and contact details are often in
        # Word headers/footers rather than the document body. Linked sections
        # expose the same part repeatedly, so index each distinct part once.
        seen_parts: set[tuple[str, str]] = set()
        for section in document.sections:
            for label, part in (("Header", section.header), ("Footer", section.footer)):
                lines = [paragraph.text for paragraph in part.paragraphs if paragraph.text.strip()]
                lines.extend(" | ".join(cell.text for cell in row.cells)
                             for table in part.tables for row in table.rows)
                text = "\n".join(line for line in lines if line.strip())
                if text and (label, text) not in seen_parts:
                    paragraphs.append(f"[{label}]\n{text}")
                    seen_parts.add((label, text))
        return ExtractionResult("\n".join(paragraphs), {"tables": len(document.tables)})

    def _xlsx(self, content: bytes) -> ExtractionResult:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        cached = None
        output: list[str] = []
        sheet_names: list[str] = []
        try:
            cached = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            for worksheet, cached_sheet in zip(workbook.worksheets, cached.worksheets, strict=True):
                sheet_names.append(worksheet.title)
                output.append(f"[Sheet: {worksheet.title}]")
                for row, cached_row in zip(worksheet.iter_rows(values_only=True),
                                           cached_sheet.iter_rows(values_only=True), strict=True):
                    values = []
                    for value, cached_value in zip(row, cached_row, strict=True):
                        if isinstance(value, str) and value.startswith("=") and cached_value is not None:
                            values.append(f"{cached_value} [cached formula result: {value}]")
                        else:
                            values.append("" if value is None else str(value))
                    if any(values):
                        output.append(" | ".join(values))
        finally:
            workbook.close()
            if cached is not None:
                cached.close()
        return ExtractionResult("\n".join(output), {"sheets": sheet_names})

    def _pptx(self, content: bytes) -> ExtractionResult:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        presentation = Presentation(io.BytesIO(content))
        output: list[str] = []

        def shape_lines(shape: Any) -> list[str]:
            lines = []
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                lines.append(shape.text)
            if getattr(shape, "has_table", False):
                lines.extend(" | ".join(cell.text for cell in row.cells)
                             for row in shape.table.rows)
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                for child in shape.shapes:
                    lines.extend(shape_lines(child))
            return lines

        for index, slide in enumerate(presentation.slides, start=1):
            output.append(f"[Slide {index}]")
            for shape in slide.shapes:
                output.extend(shape_lines(shape))
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame
                if notes is not None and notes.text.strip():
                    output.append(f"[Slide {index} Notes]\n{notes.text}")
        return ExtractionResult("\n".join(output), {"slides": len(presentation.slides)})

    def _image_ocr(self, content: bytes, suffix: str) -> ExtractionResult:
        try:
            import Vision
            from Foundation import NSURL
        except ImportError:
            return ExtractionResult(
                "", warnings=["Install pyobjc-framework-Vision to enable local OCR"]
            )

        with tempfile.TemporaryDirectory(prefix="email-reader-ocr-") as tmp:
            image_path = Path(tmp) / f"image{suffix}"
            image_path.write_bytes(content)
            request = Vision.VNRecognizeTextRequest.alloc().init()
            request.setRecognitionLevel_(Vision.VNRequestTextRecognitionLevelAccurate)
            handler = Vision.VNImageRequestHandler.alloc().initWithURL_options_(
                NSURL.fileURLWithPath_(str(image_path)), None
            )
            success, error = handler.performRequests_error_([request], None)
            if not success:
                return ExtractionResult("", warnings=[f"Vision OCR failed: {error}"])
            lines: list[str] = []
            for observation in request.results() or []:
                candidates = observation.topCandidates_(1)
                if candidates:
                    lines.append(str(candidates[0].string()))
            return ExtractionResult("\n".join(lines), {"ocr": "macOS Vision"})
