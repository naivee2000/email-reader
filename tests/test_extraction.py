from __future__ import annotations

import io
from unittest.mock import Mock
from zipfile import ZipFile

import fitz
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation

from email_reader.extraction import AttachmentExtractor, ExtractionResult


def test_pdf_keeps_native_header_and_ocrs_large_scanned_body():
    image = io.BytesIO()
    Image.new("RGB", (800, 800), "white").save(image, format="PNG")
    with fitz.open() as document:
        page = document.new_page()
        page.insert_text((72, 72), "Selectable account summary header longer than thirty characters.")
        page.insert_image(fitz.Rect(40, 110, 550, 780), stream=image.getvalue())
        content = document.tobytes()
    extractor = AttachmentExtractor()
    extractor._image_ocr = Mock(return_value=ExtractionResult(
        "Selectable account summary header longer than thirty characters.\n"
        "Scanned policy renewal date November 9, 2026.", {"ocr": "macOS Vision"}
    ))
    result = extractor.extract("mixed-page.pdf", "application/pdf", content)
    assert "Selectable account summary header" in result.text
    assert "Scanned policy renewal date November 9, 2026." in result.text
    assert result.text.count("Selectable account summary header") == 1
    assert result.metadata["ocr_pages"] == [1]
    extractor._image_ocr.assert_called_once()


def test_pdf_does_not_ocr_small_logo_when_native_text_is_present():
    image = io.BytesIO()
    Image.new("RGB", (80, 80), "white").save(image, format="PNG")
    with fitz.open() as document:
        page = document.new_page()
        page.insert_text((72, 72), "Selectable account summary header longer than thirty characters.")
        page.insert_image(fitz.Rect(500, 20, 535, 55), stream=image.getvalue())
        content = document.tobytes()
    extractor = AttachmentExtractor()
    extractor._image_ocr = Mock(return_value=ExtractionResult("unneeded OCR"))
    result = extractor.extract("native-with-logo.pdf", "application/pdf", content)
    assert "Selectable account summary header" in result.text
    assert "unneeded OCR" not in result.text
    extractor._image_ocr.assert_not_called()


def test_pdf_preserves_ocr_failure_warnings_and_other_pages():
    extractor = AttachmentExtractor()
    with fitz.open() as document:
        page = document.new_page()
        page.insert_text((72, 72), "Original supported deadline is September 12, 2026.")
        document.new_page()
        content = document.tobytes()
    extractor._image_ocr = Mock(return_value=ExtractionResult("", warnings=["OCR unavailable"]))
    result = extractor.extract("scanned.pdf", "application/pdf", content)
    assert "September 12, 2026" in result.text
    assert result.warnings == ["Page 2: OCR unavailable"]
    assert result.metadata["ocr_pages"] == []


def test_pdf_ocr_exception_preserves_native_text_and_marks_page_retryable():
    extractor = AttachmentExtractor()
    with fitz.open() as document:
        page = document.new_page()
        page.insert_text((72, 72), "Native text remains available despite another page failure.")
        document.new_page()
        content = document.tobytes()
    extractor._image_ocr = Mock(side_effect=RuntimeError("private detail not retained"))
    result = extractor.extract("mixed.pdf", "application/pdf", content)
    assert "Native text remains available" in result.text
    assert result.warnings == ["Page 2: OCR failed: RuntimeError"]


def test_pdf_successful_empty_ocr_has_no_failure_warning():
    extractor = AttachmentExtractor()
    with fitz.open() as document:
        document.new_page()
        content = document.tobytes()
    extractor._image_ocr = Mock(return_value=ExtractionResult("", {"ocr": "macOS Vision"}))
    result = extractor.extract("blank.pdf", "application/pdf", content)
    assert result.text == "[Page 1]"
    assert result.warnings == []


def test_extracts_text_csv_and_pdf() -> None:
    extractor = AttachmentExtractor()
    assert "hello" in extractor.extract("note.txt", "text/plain", b"hello world").text
    csv_result = extractor.extract("amounts.csv", "text/csv", b"name,amount\nRenewal,48000")
    assert "Renewal | 48000" in csv_result.text

    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "PDF renewal deadline September 12")
    pdf_result = extractor.extract("proposal.pdf", "application/pdf", document.tobytes())
    assert "September 12" in pdf_result.text
    assert pdf_result.metadata["pages"] == 1


def test_utf16_text_csv_and_tsv_are_searchable_without_nul_placeholders() -> None:
    extractor = AttachmentExtractor()
    note = extractor.extract(
        "policy.txt", "text/plain; charset=utf-16",
        "保险生效日期 2026-11-09".encode("utf-16"),
    )
    assert note.text == "保险生效日期 2026-11-09"
    assert "\x00" not in note.text

    csv_result = extractor.extract(
        "coverage.csv", "text/csv", "Coverage,Amount\nCollision,48000".encode("utf-16")
    )
    assert "Collision | 48000" in csv_result.text

    tsv_result = extractor.extract(
        "options.tsv", "text/tab-separated-values",
        "Option\tPrice\nFirst\t48000".encode("utf-16-le"),
    )
    assert "First | 48000" in tsv_result.text
    assert extractor.version_for("coverage.csv", "text/csv") == extractor.TEXT_CACHE_VERSION
    assert extractor.version_for("scan.pdf", "application/pdf") == extractor.CACHE_VERSION


def test_utf8_text_is_not_reinterpreted_as_utf16() -> None:
    result = AttachmentExtractor().extract(
        "note.txt", "text/plain", "Résumé and renewal date".encode()
    )
    assert result.text == "Résumé and renewal date"


def test_extracts_word_excel_and_powerpoint() -> None:
    extractor = AttachmentExtractor()

    word = Document()
    word.add_paragraph("Word contract value is 48000")
    word_bytes = io.BytesIO()
    word.save(word_bytes)
    assert "contract value" in extractor.extract("contract.docx", "", word_bytes.getvalue()).text

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Options"
    worksheet.append(["Option", "Price"])
    worksheet.append([1, 48000])
    excel_bytes = io.BytesIO()
    workbook.save(excel_bytes)
    excel = extractor.extract("options.xlsx", "", excel_bytes.getvalue())
    assert "[Sheet: Options]" in excel.text
    assert "1 | 48000" in excel.text

    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = "Renewal decision"
    powerpoint_bytes = io.BytesIO()
    slides.save(powerpoint_bytes)
    presentation = extractor.extract("review.pptx", "", powerpoint_bytes.getvalue())
    assert "[Slide 1]" in presentation.text
    assert "Renewal decision" in presentation.text


def test_word_headers_and_footers_are_indexed_once_across_linked_sections() -> None:
    document = Document()
    document.add_paragraph("Body details")
    document.sections[0].header.paragraphs[0].text = "Policy number BMW-2026"
    document.sections[0].footer.paragraphs[0].text = "Renewal November 9, 2026"
    document.add_section()
    serialized = io.BytesIO()
    document.save(serialized)

    result = AttachmentExtractor().extract("policy.docx", "", serialized.getvalue())
    assert result.text.count("Policy number BMW-2026") == 1
    assert result.text.count("Renewal November 9, 2026") == 1
    assert AttachmentExtractor.version_for("policy.docx", "") == 3


def test_excel_indexes_saved_calculated_value_and_labels_formula_cache() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Premium"
    sheet.append(["Base", "Total"])
    sheet.append([48000, "=A2+4500"])
    serialized = io.BytesIO()
    workbook.save(serialized)
    with ZipFile(io.BytesIO(serialized.getvalue())) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    original = b"<f>A2+4500</f><v></v>"
    assert original in entries["xl/worksheets/sheet1.xml"]
    entries["xl/worksheets/sheet1.xml"] = entries["xl/worksheets/sheet1.xml"].replace(
        original, b"<f>A2+4500</f><v>52500</v>"
    )
    calculated = io.BytesIO()
    with ZipFile(calculated, "w") as archive:
        for name, data in entries.items():
            archive.writestr(name, data)

    extractor = AttachmentExtractor()
    result = extractor.extract("premium.xlsx", "", calculated.getvalue())
    assert "48000 | 52500 [cached formula result: =A2+4500]" in result.text
    without_cache = extractor.extract("premium.xlsx", "", serialized.getvalue())
    assert "48000 | =A2+4500" in without_cache.text
    assert "cached formula result" not in without_cache.text
    assert extractor.version_for("premium.xlsx", "") == 3


def test_powerpoint_indexes_grouped_text_tables_and_speaker_notes() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(0, 0, 1_000_000, 200_000).text = "Grouped renewal decision"
    table = slide.shapes.add_table(2, 2, 0, 300_000, 2_000_000, 600_000).table
    table.cell(0, 0).text = "Option"
    table.cell(0, 1).text = "Premium"
    table.cell(1, 0).text = "Full coverage"
    table.cell(1, 1).text = "52500"
    slide.notes_slide.notes_text_frame.text = "Confirm before November 9, 2026"
    serialized = io.BytesIO()
    presentation.save(serialized)

    extractor = AttachmentExtractor()
    result = extractor.extract("renewal.pptx", "", serialized.getvalue())
    assert "Grouped renewal decision" in result.text
    assert "Full coverage | 52500" in result.text
    assert "[Slide 1 Notes]\nConfirm before November 9, 2026" in result.text
    assert extractor.version_for("renewal.pptx", "") == 3
